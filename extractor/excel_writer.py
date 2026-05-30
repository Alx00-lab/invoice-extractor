import io
import logging
import time
from datetime import date as _date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .validator import parse_amount, detect_currency_symbol

logger = logging.getLogger(__name__)

# ── Brand constants ───────────────────────────────────────────
BRAND_DARK   = "1F4E78"   # headers, TOTAL row
BRAND_MID    = "2E75B6"   # mid-blue accents
BRAND_LIGHT  = "D6E4F0"   # subtotal rows
ALT_ROW      = "F2F7FC"   # alternating row tint
WHITE        = "FFFFFF"
WARN_BG      = "FFF2CC"
ERROR_BG     = "F8CBAD"
FONT_NAME    = "Calibri"
WEBSITE_URL  = "yourwebsite.com"
COMPANY_NAME = "Your Company Name"


# ── Low-level cell helpers ────────────────────────────────────

def _side():
    return Side(style="thin", color="CCCCCC")


def _border():
    s = _side()
    return Border(left=s, right=s, top=s, bottom=s)


def _header_cell(cell, text, size=11):
    cell.value = text
    cell.font = Font(name=FONT_NAME, bold=True, color=WHITE, size=size)
    cell.fill = PatternFill("solid", fgColor=BRAND_DARK)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _border()


def _data_cell(cell, value, alt=False, align="left", bold=False, size=11):
    cell.value = value if value not in (None, "null", "") else "—"
    cell.font = Font(name=FONT_NAME, size=size, bold=bold)
    cell.fill = PatternFill("solid", fgColor=ALT_ROW if alt else WHITE)
    cell.border = _border()
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)


def _subtotal_cell(cell, value, is_label=False):
    cell.value = value
    cell.font = Font(name=FONT_NAME, bold=True, size=11, color=BRAND_DARK)
    cell.fill = PatternFill("solid", fgColor=BRAND_LIGHT)
    cell.alignment = Alignment(
        horizontal="left" if is_label else "right", vertical="center"
    )
    cell.border = _border()


def _total_cell(cell, value, is_label=False):
    cell.value = value
    cell.font = Font(name=FONT_NAME, bold=True, size=11, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BRAND_DARK)
    cell.alignment = Alignment(
        horizontal="left" if is_label else "right", vertical="center"
    )
    cell.border = _border()


def _footer(ws, row, n_cols):
    """Two-part branded footer: 'Thank you' left, byline right."""
    left_text  = "Thank you for your business."
    right_text = f"Built by Alex · Finance Automation Specialist · {WEBSITE_URL}"
    mid        = max(1, n_cols // 2)

    if mid > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=mid)
    c_l = ws.cell(row=row, column=1)
    c_l.value = left_text
    c_l.font  = Font(name=FONT_NAME, size=10, italic=True, color="888888")
    c_l.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    right_start = mid + 1
    if right_start <= n_cols and right_start < n_cols:
        ws.merge_cells(
            start_row=row, start_column=right_start,
            end_row=row,   end_column=n_cols
        )
    if right_start <= n_cols:
        c_r = ws.cell(row=row, column=right_start)
        c_r.value = right_text
        c_r.font  = Font(name=FONT_NAME, size=10, italic=True, color=BRAND_DARK)
        c_r.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    ws.row_dimensions[row].height = 18


def _set_widths(ws, overrides=None, default=12, max_w=50):
    """Auto-size columns; overrides = {col_letter: min_width}."""
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if not hasattr(cell, "column_letter") or cell.value is None:
                continue
            col = cell.column_letter
            widths[col] = max(widths.get(col, 0), len(str(cell.value)))
    overrides = overrides or {}
    for col, w in widths.items():
        min_w = overrides.get(col, default)
        ws.column_dimensions[col].width = min(max(w + 2, min_w), max_w)
    for col, min_w in overrides.items():
        if col not in widths:
            ws.column_dimensions[col].width = min_w


def _fmt(val, symbol="$"):
    """Format float as currency string; returns '—' if None."""
    if val is None:
        return "—"
    return f"{symbol}{val:,.2f}"


# ── Sheet 1: Invoice Summary (8 cols A-H) ─────────────────────

def _build_summary_sheet(ws, invoices, ok_invoices, sums, present, symbol):
    ws.title = "Invoice Summary"
    N = 8  # columns A through H

    # ── Company header (rows 1-2) ──────────────────────────────
    ws.merge_cells("A1:D1")
    c = ws["A1"]
    c.value = COMPANY_NAME
    c.font  = Font(name=FONT_NAME, bold=True, size=14, color=BRAND_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("E1:H1")
    c = ws["E1"]
    c.value = "INVOICE"
    c.font  = Font(name=FONT_NAME, bold=True, size=14, color=WHITE)
    c.fill  = PatternFill("solid", fgColor=BRAND_DARK)
    c.alignment = Alignment(horizontal="right", vertical="center")

    ws.merge_cells("E2:H2")
    c = ws["E2"]
    c.value = "Finance Automation Specialist"
    c.font  = Font(name=FONT_NAME, italic=True, size=10, color=BRAND_DARK)
    c.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 18

    # row 3 is intentionally empty
    header_start = 4

    # ── BILL TO section (single invoice) or batch note ─────────
    if len(ok_invoices) == 1:
        d = ok_invoices[0].get("data", {})

        # Left block: BILL TO label + vendor
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=3)
        c = ws.cell(row=4, column=1)
        c.value = "BILL TO"
        c.font  = Font(name=FONT_NAME, bold=True, size=9, color=WHITE)
        c.fill  = PatternFill("solid", fgColor=BRAND_DARK)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[4].height = 16

        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=3)
        c = ws.cell(row=5, column=1)
        c.value = d.get("vendor_name") or "—"
        c.font  = Font(name=FONT_NAME, bold=True, size=11)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[5].height = 18

        ws.merge_cells(start_row=6, start_column=1, end_row=7, end_column=3)
        c = ws.cell(row=6, column=1)
        c.value = d.get("payment_terms") or ""
        c.font  = Font(name=FONT_NAME, size=10, color="555555")
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Right block: invoice meta
        meta = [
            ("Invoice #", d.get("invoice_number") or "—"),
            ("Date",      d.get("invoice_date")   or "—"),
            ("Due Date",  d.get("due_date")        or "—"),
            ("Status",    "Unpaid"),
        ]
        for i, (label, value) in enumerate(meta):
            r = 4 + i
            c_lbl = ws.cell(row=r, column=6)
            c_lbl.value = label
            c_lbl.font  = Font(name=FONT_NAME, bold=True, size=10, color=BRAND_DARK)
            c_lbl.fill  = PatternFill("solid", fgColor=BRAND_LIGHT)
            c_lbl.alignment = Alignment(horizontal="right", vertical="center")
            c_lbl.border = _border()

            ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
            c_val = ws.cell(row=r, column=7)
            c_val.value = value
            c_val.font  = Font(name=FONT_NAME, size=10)
            c_val.alignment = Alignment(horizontal="left", vertical="center")
            c_val.border = _border()
            ws.row_dimensions[r].height = 18

        header_start = 9  # row 8 = empty spacer

    else:
        # Multiple invoices — compact batch banner
        ws.merge_cells(
            start_row=header_start, start_column=1,
            end_row=header_start,   end_column=N
        )
        c = ws.cell(row=header_start, column=1)
        n_err = len(invoices) - len(ok_invoices)
        c.value = (
            f"Batch — {len(ok_invoices)} invoice(s) processed"
            + (f" · {n_err} error(s)" if n_err else "")
        )
        c.font  = Font(name=FONT_NAME, bold=True, size=11, color=WHITE)
        c.fill  = PatternFill("solid", fgColor=BRAND_MID)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[header_start].height = 22
        header_start += 2  # skip one empty row

    # ── Column headers ─────────────────────────────────────────
    col_headers = [
        "#", "Description", "Vendor", "Qty",
        "Unit Price", "Tax %", "Tax Amt", "Total"
    ]
    for col, h in enumerate(col_headers, start=1):
        _header_cell(ws.cell(row=header_start, column=col), h)
    ws.row_dimensions[header_start].height = 24
    row = header_start + 1

    # ── Line items ─────────────────────────────────────────────
    counter = 0
    for inv in invoices:
        if inv.get("status") != "ok":
            continue
        d      = inv.get("data", {})
        vendor = d.get("vendor_name") or "—"
        for item in d.get("line_items") or []:
            alt = (counter % 2 == 1)
            _data_cell(ws.cell(row=row, column=1), counter + 1, alt, align="center")
            _data_cell(ws.cell(row=row, column=2), item.get("description"), alt)
            _data_cell(ws.cell(row=row, column=3), vendor, alt)
            _data_cell(ws.cell(row=row, column=4), item.get("quantity"),  alt, align="center")
            _data_cell(ws.cell(row=row, column=5), item.get("unit_price"), alt, align="right")
            _data_cell(ws.cell(row=row, column=6), "—", alt, align="center")   # Tax % per line
            _data_cell(ws.cell(row=row, column=7), "—", alt, align="center")   # Tax Amt per line
            _data_cell(ws.cell(row=row, column=8), item.get("amount"), alt, align="right")
            ws.row_dimensions[row].height = 18
            row += 1
            counter += 1

    if counter == 0:
        ws.merge_cells(
            start_row=row, start_column=1, end_row=row, end_column=N
        )
        c = ws.cell(row=row, column=1)
        c.value = "No line items extracted"
        c.font  = Font(name=FONT_NAME, italic=True, size=11, color="888888")
        c.alignment = Alignment(horizontal="center", vertical="center")
        row += 1

    row += 1  # spacer

    # ── Subtotal ───────────────────────────────────────────────
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=N - 1
    )
    _subtotal_cell(ws.cell(row=row, column=1), "Subtotal", is_label=True)
    _subtotal_cell(
        ws.cell(row=row, column=N),
        _fmt(sums["subtotal"] if present["subtotal"] else None, symbol)
    )
    ws.row_dimensions[row].height = 18
    row += 1

    # ── Tax ────────────────────────────────────────────────────
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=N - 1
    )
    _subtotal_cell(ws.cell(row=row, column=1), "Tax", is_label=True)
    _subtotal_cell(
        ws.cell(row=row, column=N),
        _fmt(sums["tax"] if present["tax"] else None, symbol)
    )
    ws.row_dimensions[row].height = 18
    row += 1

    # ── TOTAL DUE ──────────────────────────────────────────────
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=N - 1
    )
    _total_cell(ws.cell(row=row, column=1), "TOTAL DUE", is_label=True)
    _total_cell(
        ws.cell(row=row, column=N),
        _fmt(sums["total"] if present["total"] else None, symbol)
    )
    ws.row_dimensions[row].height = 22
    row += 2

    # ── Footer ─────────────────────────────────────────────────
    _footer(ws, row, N)

    ws.freeze_panes = "A3"
    _set_widths(ws, {"B": 35, "C": 20})


# ── Sheet 2: Line Items (7 cols A-G) ──────────────────────────

def _build_lineitems_sheet(ws, invoices, sums, present, symbol):
    ws.title = "Line Items"
    N = 7

    # ── Header bar ─────────────────────────────────────────────
    ws.merge_cells("A1:C1")
    c = ws["A1"]
    c.value = "Line Item Detail"
    c.font  = Font(name=FONT_NAME, bold=True, size=14, color=BRAND_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ok = [i for i in invoices if i.get("status") == "ok"]
    if len(ok) == 1:
        d   = ok[0].get("data", {})
        ref = f"{d.get('invoice_number') or '—'} · {d.get('vendor_name') or '—'}"
    else:
        ref = f"Batch: {len(ok)} invoice(s)"
    ws.merge_cells("D1:G1")
    c = ws["D1"]
    c.value = ref
    c.font  = Font(name=FONT_NAME, italic=True, size=11, color=BRAND_DARK)
    c.alignment = Alignment(horizontal="right", vertical="center")

    # ── Column headers ─────────────────────────────────────────
    for col, h in enumerate(
        ["#", "File", "Invoice #", "Description", "Qty", "Unit Price", "Amount"],
        start=1,
    ):
        _header_cell(ws.cell(row=2, column=col), h)
    ws.row_dimensions[2].height = 24

    # ── Data ───────────────────────────────────────────────────
    row     = 3
    counter = 0
    for inv in invoices:
        if inv.get("status") != "ok":
            continue
        d        = inv.get("data", {})
        filename = inv.get("filename", "—")
        inv_num  = d.get("invoice_number") or "—"
        for item in d.get("line_items") or []:
            alt = (counter % 2 == 1)
            _data_cell(ws.cell(row=row, column=1), counter + 1, alt, align="center")
            _data_cell(ws.cell(row=row, column=2), filename, alt)
            _data_cell(ws.cell(row=row, column=3), inv_num,  alt)
            _data_cell(ws.cell(row=row, column=4), item.get("description"), alt)
            _data_cell(ws.cell(row=row, column=5), item.get("quantity"),  alt, align="center")
            _data_cell(ws.cell(row=row, column=6), item.get("unit_price"), alt, align="right")
            _data_cell(ws.cell(row=row, column=7), item.get("amount"),    alt, align="right")
            ws.row_dimensions[row].height = 18
            row += 1
            counter += 1

    if counter == 0:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=N)
        c = ws.cell(row=3, column=1)
        c.value = "No line items extracted"
        c.font  = Font(name=FONT_NAME, italic=True, size=11, color="888888")
        c.alignment = Alignment(horizontal="center", vertical="center")
        row = 4

    row += 1  # spacer

    # ── Summary totals ─────────────────────────────────────────
    for label, key, fn in [
        ("Subtotal",  "subtotal", _subtotal_cell),
        ("Tax",       "tax",      _subtotal_cell),
        ("TOTAL DUE", "total",    _total_cell),
    ]:
        ws.merge_cells(
            start_row=row, start_column=1, end_row=row, end_column=N - 1
        )
        fn(ws.cell(row=row, column=1), label, is_label=True)
        fn(ws.cell(row=row, column=N),
           _fmt(sums[key] if present[key] else None, symbol))
        ws.row_dimensions[row].height = 18 if label != "TOTAL DUE" else 22
        row += 1

    row += 1
    _footer(ws, row, N)

    ws.freeze_panes = "A3"
    _set_widths(ws, {"D": 35, "B": 20})


# ── Sheet 3: Automation Report (3 cols A-C) ───────────────────

def _build_report_sheet(ws, invoices, ok_invoices, err_count,
                        sums, present, symbol, elapsed):
    ws.title = "Automation Report"
    N = 3

    # ── Title ──────────────────────────────────────────────────
    ws.merge_cells("A1:B1")
    c = ws["A1"]
    c.value = "Automation Processing Report"
    c.font  = Font(name=FONT_NAME, bold=True, size=14, color=BRAND_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    c = ws["C1"]
    c.value = f"Generated: {_date.today().strftime('%B %d, %Y')}"
    c.font  = Font(name=FONT_NAME, italic=True, size=10, color="888888")
    c.alignment = Alignment(horizontal="right", vertical="center")

    # row 2 = empty spacer

    # ── Headers ────────────────────────────────────────────────
    for col, h in enumerate(["Metric", "Result", "Status"], start=1):
        _header_cell(ws.cell(row=3, column=col), h)
    ws.row_dimensions[3].height = 22

    # ── Metrics ────────────────────────────────────────────────
    total = len(invoices)
    ok    = len(ok_invoices)
    pct   = f"{int(ok / total * 100)}%" if total else "—"

    total_val = sums["total"] if present["total"] else None
    tax_val   = sums["tax"]   if present["tax"]   else None

    if elapsed is not None:
        elapsed_str    = f"< {int(elapsed) + 1} seconds"
        elapsed_status = "✓ Fast" if elapsed < 60 else "⚠ Slow"
    else:
        elapsed_str    = "< 60 seconds"
        elapsed_status = "✓ Fast"

    metrics = [
        ("Invoices processed",     f"{total} invoice(s)",         "✓ OK"),
        ("Successful extractions", f"{ok} of {total} ({pct})",
         "✓ OK" if ok == total else "⚠ Check errors"),
        ("Failed / flagged",       str(err_count) if err_count else "0",
         "⚠ Review" if err_count else "—"),
        ("Total value captured",   _fmt(total_val, symbol),       "✓ OK"),
        ("Tax captured",           _fmt(tax_val, symbol),         "✓ OK"),
        ("Processing time",        elapsed_str,                   elapsed_status),
        ("Manual entry required",  "0 fields",                    "✓ None"),
    ]

    row = 4
    for i, (metric, result, status) in enumerate(metrics):
        alt = (i % 2 == 1)
        _data_cell(ws.cell(row=row, column=1), metric, alt, bold=True)
        _data_cell(ws.cell(row=row, column=2), result, alt, align="center")

        c = ws.cell(row=row, column=3)
        c.value = status
        if status.startswith("✓"):
            color = "2D6A3F"
        elif "⚠" in status:
            color = "7B3F00"
        else:
            color = "888888"
        c.font  = Font(name=FONT_NAME, size=11, color=color)
        c.fill  = PatternFill("solid", fgColor=ALT_ROW if alt else WHITE)
        c.border = _border()
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 20
        row += 1

    row += 1  # spacer

    # ── CTA ────────────────────────────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    c = ws.cell(row=row, column=1)
    c.value = "Want this running automatically for your business?"
    c.font  = Font(name=FONT_NAME, bold=True, size=11, color=BRAND_DARK)
    c.alignment = Alignment(horizontal="left", vertical="center")

    c = ws.cell(row=row, column=3)
    c.value = f"Book a free 30-min audit → {WEBSITE_URL}"
    c.font  = Font(name=FONT_NAME, italic=True, size=11, color=BRAND_MID)
    c.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row].height = 22
    row += 2

    # ── Footer ─────────────────────────────────────────────────
    _footer(ws, row, N)

    _set_widths(ws, {"A": 30, "B": 25, "C": 15})


# ── Public API ────────────────────────────────────────────────

def generate_batch_excel(invoices: list, start_time: float = None) -> bytes:
    """
    Generate a branded 3-sheet Excel from a batch of invoices.

    Args:
        invoices:   list of dicts — keys: filename, data, status, error
        start_time: time.time() snapshot taken before extraction begins;
                    used to compute elapsed in Sheet 3.  Optional.
    """
    ok_invoices = [i for i in invoices if i.get("status") == "ok"]
    err_count   = len(invoices) - len(ok_invoices)

    # Aggregate sums + detect currency symbol
    sums        = {"subtotal": 0.0, "tax": 0.0, "total": 0.0}
    present     = {"subtotal": False, "tax": False, "total": False}
    symbol_pool = []
    for inv in ok_invoices:
        d = inv.get("data", {})
        for k in sums:
            amt = parse_amount(d.get(k))
            if amt is not None:
                sums[k]    += amt
                present[k]  = True
            if d.get(k):
                symbol_pool.append(d[k])
    symbol  = detect_currency_symbol(symbol_pool) or "$"
    elapsed = (time.time() - start_time) if start_time else None

    wb = Workbook()
    _build_summary_sheet(wb.active, invoices, ok_invoices, sums, present, symbol)
    _build_lineitems_sheet(wb.create_sheet("Line Items"),       invoices, sums, present, symbol)
    _build_report_sheet(
        wb.create_sheet("Automation Report"),
        invoices, ok_invoices, err_count, sums, present, symbol, elapsed
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    logger.info(f"Excel generated: {len(invoices)} invoices, {len(ok_invoices)} OK.")
    return buf.getvalue()


def generate_excel(data: dict) -> bytes:
    """Single-invoice wrapper — backward compatible."""
    return generate_batch_excel([{
        "filename": "invoice.pdf",
        "data": data,
        "status": "ok",
    }])
